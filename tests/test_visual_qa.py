"""Tests for pptx_qa.py using a stubbed `pptx` module.

python-pptx can't be installed in every environment, so these tests inject a
minimal fake that implements the small surface pptx_qa actually uses
(slides, shapes, text frames, fonts, EMU dimensions). This exercises the real
audit logic: slide count, font census, overflow heuristic, placeholder scan,
and the empty-slide check.
"""

import contextlib
import io
import sys
import types
import unittest
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent


def _install_pptx_stub():
    if "pptx" in sys.modules:
        return
    pptx = types.ModuleType("pptx")
    enum_mod = types.ModuleType("pptx.enum")
    shapes_mod = types.ModuleType("pptx.enum.shapes")

    class MSO_SHAPE_TYPE:
        PICTURE = 13
        AUTO_SHAPE = 1

    shapes_mod.MSO_SHAPE_TYPE = MSO_SHAPE_TYPE
    enum_mod.shapes = shapes_mod
    pptx.Presentation = None  # patched per-test via pptx_qa.Presentation
    sys.modules["pptx"] = pptx
    sys.modules["pptx.enum"] = enum_mod
    sys.modules["pptx.enum.shapes"] = shapes_mod


_install_pptx_stub()

sys.path.insert(0, str(REPO / "visual-qa"))
import pptx_qa  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE_TYPE  # noqa: E402

EMU_PER_INCH = 914400


class _Font:
    def __init__(self, name):
        self.name = name


class _Run:
    def __init__(self, font_name):
        self.font = _Font(font_name)


class _Para:
    def __init__(self, font_name):
        self.runs = [_Run(font_name)]


class _TextFrame:
    def __init__(self, font_name):
        self.paragraphs = [_Para(font_name)]


class _Shape:
    def __init__(self, text="", width_in=6, height_in=3,
                 font_name="Calibri", picture=False, table=False):
        self._text = text
        self.width = int(width_in * EMU_PER_INCH)
        self.height = int(height_in * EMU_PER_INCH)
        self.has_text_frame = bool(text)
        self.text_frame = _TextFrame(font_name)
        self.shape_type = (MSO_SHAPE_TYPE.PICTURE if picture
                           else MSO_SHAPE_TYPE.AUTO_SHAPE)
        self.has_table = table

    @property
    def text(self):
        return self._text


class _Slide:
    def __init__(self, shapes):
        self.shapes = shapes


class _Deck:
    def __init__(self, slides, width_in=13.33, height_in=7.5):
        self.slides = slides
        self.slide_width = int(width_in * EMU_PER_INCH)
        self.slide_height = int(height_in * EMU_PER_INCH)


def messy_deck():
    return _Deck([
        _Slide([
            _Shape("Quarterly results", font_name="Calibri"),
            # 3000 chars crammed into a half-square-inch shape -> overflow
            _Shape("x" * 3000, width_in=1, height_in=0.5, font_name="Arial"),
        ]),
        _Slide([
            _Shape("Lorem ipsum dolor sit amet", font_name="Times New Roman"),
            _Shape("Agenda", font_name="Verdana"),
        ]),
        _Slide([]),  # empty slide
    ])


def clean_deck():
    return _Deck([
        _Slide([_Shape("Title", font_name="Calibri"),
                _Shape("Some reasonable body copy.", font_name="Calibri")]),
        _Slide([_Shape("Summary", font_name="Calibri")]),
    ])


def run_audit(deck, expected_slides):
    original = pptx_qa.Presentation
    pptx_qa.Presentation = lambda path: deck  # noqa: E731
    buf = io.StringIO()
    try:
        with contextlib.redirect_stdout(buf):
            code = pptx_qa.audit(Path("fake.pptx"), expected_slides)
    finally:
        pptx_qa.Presentation = original
    return code, buf.getvalue()


class TestDeckAudit(unittest.TestCase):
    def test_messy_deck_reports_every_issue(self):
        code, out = run_audit(messy_deck(), expected_slides=5)
        self.assertEqual(code, 1)
        self.assertIn("Slide count is 3, expected 5", out)
        self.assertIn("possible text overflow", out)
        self.assertIn("placeholder text", out)
        self.assertIn("appears empty", out)
        self.assertIn("More than 3 fonts", out)

    def test_clean_deck_passes(self):
        code, out = run_audit(clean_deck(), expected_slides=2)
        self.assertEqual(code, 0)
        self.assertIn("PASS", out)

    def test_overflow_heuristic_boundaries(self):
        roomy = _Shape("x" * 100, width_in=6, height_in=3)
        cramped = _Shape("x" * 3000, width_in=1, height_in=0.5)
        self.assertFalse(pptx_qa.shape_text_overflow(roomy))
        self.assertTrue(pptx_qa.shape_text_overflow(cramped))
        self.assertFalse(pptx_qa.shape_text_overflow(_Shape()))


if __name__ == "__main__":
    unittest.main()
